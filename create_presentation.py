# -*- coding: utf-8 -*-
"""Create 6-slide PPTX presentation for NIR project."""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

OUT = Path(r"C:\Users\janap\Downloads\НИРМ_Презентация.pptx")
SHOTS = Path(r"C:\Users\janap\Downloads\resume-ai-demo\screenshots")

# Colors
BG = RGBColor(15, 23, 42)       # #0f172a
ACCENT = RGBColor(14, 165, 233)  # #0ea5e9
WHITE = RGBColor(226, 232, 240)
GRAY = RGBColor(148, 163, 184)
GREEN = RGBColor(34, 197, 94)


def set_bg(slide, color=BG):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_textbox(slide, left, top, width, height, text, size=18, bold=False, color=WHITE, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = color
    p.font.name = "Segoe UI"
    p.alignment = align
    return box


def add_bullets(slide, left, top, width, height, items, size=16, color=WHITE):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = item
        p.font.size = Pt(size)
        p.font.color.rgb = color
        p.font.name = "Segoe UI"
        p.space_after = Pt(8)
        p.level = 0
    return box


def slide_header(slide, title, subtitle=None):
    set_bg(slide)
    add_textbox(slide, Inches(0.6), Inches(0.35), Inches(8.5), Inches(0.7), title, size=28, bold=True, color=ACCENT)
    if subtitle:
        add_textbox(slide, Inches(0.6), Inches(0.95), Inches(8.5), Inches(0.5), subtitle, size=14, color=GRAY)
    # accent line
    line = slide.shapes.add_shape(1, Inches(0.6), Inches(1.35), Inches(1.2), Inches(0.04))
    line.fill.solid()
    line.fill.fore_color.rgb = ACCENT
    line.line.fill.background()


def main():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    # ── SLIDE 1: Title ──
    s1 = prs.slides.add_slide(blank)
    set_bg(s1)
    add_textbox(s1, Inches(1), Inches(2.0), Inches(11), Inches(1.2),
                "Resume AI", size=44, bold=True, color=ACCENT, align=PP_ALIGN.CENTER)
    add_textbox(s1, Inches(1), Inches(3.1), Inches(11), Inches(0.8),
                "Жасанды интеллект негізіндегі түйіндеме веб-қосымшасы", size=22, color=WHITE, align=PP_ALIGN.CENTER)
    add_textbox(s1, Inches(1), Inches(4.2), Inches(11), Inches(0.6),
                "НИРМ — жобалық зерттеу жұмысы", size=16, color=GRAY, align=PP_ALIGN.CENTER)
    add_textbox(s1, Inches(1), Inches(5.5), Inches(11), Inches(0.5),
                "Орындаушы: Ырысбек  ·  2026", size=14, color=GRAY, align=PP_ALIGN.CENTER)

    # ── SLIDE 2: Problem & Goal ──
    s2 = prs.slides.add_slide(blank)
    slide_header(s2, "Мәселе және мақсат")
    add_bullets(s2, Inches(0.6), Inches(1.7), Inches(5.8), Inches(5.0), [
        "Жұмыс іздеушілер түйіндеме дайындауға көп уақыт жұмсауда",
        "Құжаттар жұмыс орны талаптарына сәйкес келмейді",
        "Жұмыс берушілер сапасыз өтінімдермен шамадан тыс жүктеледі",
        "Колданыстағы шешімдер терең AI талдауын қамтамасыз етпейді",
    ], size=17)
    add_textbox(s2, Inches(6.8), Inches(1.7), Inches(5.8), Inches(0.5),
                "Жоба мақсаты", size=20, bold=True, color=GREEN)
    add_bullets(s2, Inches(6.8), Inches(2.3), Inches(5.8), Inches(4.5), [
        "AI модельдерін біріктіретін веб-қосымша әзірлеу",
        "Түйіндемені автоматты құру және талдау",
        "80%+ талдау дәлдігін қамтамасыз ету",
        "Жұмыс орны талаптарына бейімдеу",
    ], size=17, color=RGBColor(200, 220, 240))

    # ── SLIDE 3: Technologies & Architecture ──
    s3 = prs.slides.add_slide(blank)
    slide_header(s3, "Технологиялар және архитектура")
    add_textbox(s3, Inches(0.6), Inches(1.7), Inches(5.5), Inches(0.4),
                "Технологиялық стек", size=18, bold=True, color=ACCENT)
    add_bullets(s3, Inches(0.6), Inches(2.2), Inches(5.5), Inches(4.5), [
        "Frontend: React — интерактивті UI, WYSIWYG редактор",
        "Backend: FastAPI — асинхронды REST API",
        "AI: BERT — семантикалық талдау",
        "AI: GPT-3.5 — мазмұн генерациясы",
        "ML: scikit-learn — ұсыныс жүйелері",
    ], size=15)
    add_textbox(s3, Inches(6.8), Inches(1.7), Inches(5.8), Inches(0.4),
                "Архитектура (3 деңгей)", size=18, bold=True, color=ACCENT)
    add_bullets(s3, Inches(6.8), Inches(2.2), Inches(5.8), Inches(4.5), [
        "Клиент — React веб-интерфейс (MVVM)",
        "Сервер — FastAPI бизнес-логика",
        "AI микросервис — NLP өңдеу (BERT + GPT)",
        "Деректер ағыны: Form → JSON → API → AI → нәтиже",
    ], size=15)

    # ── SLIDE 4: UI (with screenshot) ──
    s4 = prs.slides.add_slide(blank)
    slide_header(s4, "Пайдаланушы интерфейсі", "4 қадамдық түйіндеме құру процесі")
    img1 = SHOTS / "01-zhakt-akparat.png"
    img2 = SHOTS / "04-ai-taldau.png"
    if img1.exists():
        s4.shapes.add_picture(str(img1), Inches(0.5), Inches(1.6), width=Inches(6.0))
    if img2.exists():
        s4.shapes.add_picture(str(img2), Inches(6.8), Inches(1.6), width=Inches(6.0))
    add_textbox(s4, Inches(0.5), Inches(6.85), Inches(6), Inches(0.4),
                "1. Жеке ақпарат енгізу", size=11, color=GRAY, align=PP_ALIGN.CENTER)
    add_textbox(s4, Inches(6.8), Inches(6.85), Inches(6), Inches(0.4),
                "4. AI талдау нәтижелері", size=11, color=GRAY, align=PP_ALIGN.CENTER)

    # ── SLIDE 5: Results & Metrics ──
    s5 = prs.slides.add_slide(blank)
    slide_header(s5, "Тестілеу нәтижелері")
    metrics = [
        ("87%", "Генерация\nдәлдігі"),
        ("92%", "Precision"),
        ("85%", "Recall"),
        ("80%+", "Талдау\nдәлдігі"),
    ]
    for i, (val, label) in enumerate(metrics):
        x = Inches(0.6 + i * 3.1)
        box = s5.shapes.add_shape(1, x, Inches(2.0), Inches(2.6), Inches(1.8))
        box.fill.solid()
        box.fill.fore_color.rgb = RGBColor(30, 41, 59)
        box.line.color.rgb = ACCENT
        add_textbox(s5, x, Inches(2.15), Inches(2.6), Inches(0.9),
                    val, size=36, bold=True, color=GREEN, align=PP_ALIGN.CENTER)
        add_textbox(s5, x, Inches(3.05), Inches(2.6), Inches(0.6),
                    label, size=13, color=GRAY, align=PP_ALIGN.CENTER)

    add_bullets(s5, Inches(0.6), Inches(4.3), Inches(12), Inches(2.5), [
        "Selenium + JMeter арқылы функционалдық және жүктеме тестілеуі өткізілді",
        "AI модуль мазмұнды жақсарту бойынша тиімді кеңестер береді",
        "Жауап беру уақыты рұқсат етілген шектерде қалды",
    ], size=15)

    # ── SLIDE 6: Conclusion ──
    s6 = prs.slides.add_slide(blank)
    slide_header(s6, "Қорытынды және болашақ даму")
    add_bullets(s6, Inches(0.6), Inches(1.7), Inches(5.8), Inches(5.0), [
        "AI негізіндегі түйіндеме платформасының прототипі сәтті жасалды",
        "Түйіндеме құру + талдау + ұсыныс беру циклі іске асады",
        "Жоба мақсатына қол жеткізілді (80%+ дәлдік)",
        "HR саласын цифрландыруға практикалық үлес",
    ], size=17)
    add_textbox(s6, Inches(6.8), Inches(1.7), Inches(5.8), Inches(0.4),
                "Болашақ жоспар", size=18, bold=True, color=ACCENT)
    add_bullets(s6, Inches(6.8), Inches(2.2), Inches(5.8), Inches(4.5), [
        "LinkedIn / hh.kz интеграциясы",
        "Көптілділік (қаз / орыс / ағыл.)",
        "Түйіндеме шаблондары кітапханасын кеңейту",
        "BERT моделін локалды орналастыру (ONNX)",
    ], size=15)
    add_textbox(s6, Inches(0.6), Inches(6.5), Inches(12), Inches(0.5),
                "Рахмет за внимание!  ·  Сұрақтар?", size=20, bold=True, color=ACCENT, align=PP_ALIGN.CENTER)

    prs.save(str(OUT))
    Path(r"C:\Users\janap\Downloads\resume-ai-demo\merged\pptx_log.txt").write_text(
        f"Saved: {OUT}\nSize: {OUT.stat().st_size}", encoding="utf-8"
    )


if __name__ == "__main__":
    main()
