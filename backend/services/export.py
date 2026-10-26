"""Export resume text to PDF and DOCX using Classic / Modern / Minimal layouts."""

from __future__ import annotations

from io import BytesIO
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from fpdf import FPDF

from schemas import ResumeTemplate
from services.resume_i18n import headings_for
from services.resume_layout import ParsedResume, parse_resume_text

_FONTS_DIR = Path(__file__).resolve().parent.parent / "assets" / "fonts"
_DEJAVU = _FONTS_DIR / "DejaVuSans.ttf"
_DEJAVU_BOLD = _FONTS_DIR / "DejaVuSans-Bold.ttf"

_SKY = (14, 165, 233)
_WHITE = (255, 255, 255)
_SLATE = (15, 23, 42)
_MUTED = (100, 116, 139)


def _safe_filename(title: str, ext: str) -> str:
    base = "".join(c if c.isalnum() or c in " -_" else "_" for c in (title or "resume"))
    base = base.strip() or "resume"
    return f"{base[:80]}.{ext}"


def _make_pdf(*, margin: float = 18) -> FPDF:
    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=margin)
    pdf.add_page()
    pdf.add_font("DejaVu", fname=str(_DEJAVU))
    pdf.add_font("DejaVu", style="B", fname=str(_DEJAVU_BOLD))
    return pdf


def _pdf_text(
    pdf: FPDF,
    text: str,
    *,
    size: int = 11,
    style: str = "",
    line_h: float = 6,
    width: float | None = None,
    color: tuple[int, int, int] = _SLATE,
) -> None:
    pdf.set_text_color(*color)
    pdf.set_font("DejaVu", style=style, size=size)
    w = pdf.epw if width is None else width
    for line in (text.splitlines() or [""]):
        pdf.set_x(pdf.l_margin if width is None else pdf.get_x())
        pdf.multi_cell(w, line_h, line if line else " ")


def _pdf_heading(pdf: FPDF, title: str, *, underline: bool = True) -> None:
    pdf.ln(3)
    pdf.set_text_color(*_SLATE)
    pdf.set_font("DejaVu", style="B", size=11)
    pdf.set_x(pdf.l_margin)
    pdf.multi_cell(pdf.epw, 7, title)
    if underline:
        y = pdf.get_y()
        pdf.set_draw_color(*_SLATE)
        pdf.line(pdf.l_margin, y, pdf.l_margin + pdf.epw, y)
        pdf.ln(2)


def _pdf_sections_classic(pdf: FPDF, parsed: ParsedResume, language: str) -> None:
    headings = headings_for(language)
    if parsed.target:
        _pdf_heading(pdf, headings.target)
        _pdf_text(pdf, parsed.target)
    if parsed.experience:
        _pdf_heading(pdf, headings.experience)
        _pdf_text(pdf, parsed.experience)
    if parsed.skills:
        _pdf_heading(pdf, headings.skills)
        _pdf_text(pdf, parsed.skills)
    if parsed.education:
        _pdf_heading(pdf, headings.education)
        _pdf_text(pdf, parsed.education)


def _pdf_dump(pdf: FPDF, text: str) -> None:
    _pdf_text(pdf, text)


def resume_to_pdf(
    resume_text: str,
    *,
    template: str = "classic",
    language: str = "kk",
) -> bytes:
    parsed = parse_resume_text(resume_text)
    if template == "modern":
        pdf = _pdf_modern(parsed, language)
    elif template == "minimal":
        pdf = _pdf_minimal(parsed, language)
    else:
        pdf = _pdf_classic(parsed, language)
    out = BytesIO()
    pdf.output(out)
    return out.getvalue()


def _pdf_classic(parsed: ParsedResume, language: str) -> FPDF:
    pdf = _make_pdf(margin=18)
    name = parsed.name or (parsed.raw.splitlines()[0] if parsed.raw else "")
    _pdf_text(pdf, name, size=16, style="B", line_h=8)
    if parsed.headline:
        _pdf_text(pdf, parsed.headline, size=10, color=_MUTED, line_h=5)
    if parsed.has_sections:
        _pdf_sections_classic(pdf, parsed, language)
    else:
        pdf.ln(2)
        _pdf_dump(pdf, parsed.raw)
    return pdf


def _pdf_modern(parsed: ParsedResume, language: str) -> FPDF:
    pdf = _make_pdf(margin=16)
    headings = headings_for(language)
    header_h = 36
    pdf.set_fill_color(*_SKY)
    pdf.rect(0, 0, pdf.w, header_h, "F")
    pdf.set_text_color(*_WHITE)
    pdf.set_xy(16, 10)
    pdf.set_font("DejaVu", style="B", size=18)
    name = parsed.name or (parsed.raw.splitlines()[0] if parsed.raw else "")
    pdf.multi_cell(pdf.epw, 8, name)
    if parsed.headline:
        pdf.set_x(16)
        pdf.set_font("DejaVu", size=10)
        pdf.multi_cell(pdf.epw, 5, parsed.headline)
    pdf.set_text_color(*_SLATE)
    pdf.set_y(header_h + 8)

    if not parsed.has_sections:
        _pdf_dump(pdf, parsed.raw)
        return pdf

    gap = 8
    col_w = (pdf.epw - gap) / 2
    left_x = pdf.l_margin
    right_x = pdf.l_margin + col_w + gap
    start_y = pdf.get_y()

    def write_col(x: float, y: float, blocks: list[tuple[str, str]]) -> float:
        pdf.set_xy(x, y)
        for title, body in blocks:
            if not body:
                continue
            pdf.set_x(x)
            pdf.set_font("DejaVu", style="B", size=10)
            pdf.set_text_color(*_SKY)
            pdf.multi_cell(col_w, 6, title)
            pdf.set_x(x)
            pdf.set_font("DejaVu", size=10)
            pdf.set_text_color(*_SLATE)
            pdf.multi_cell(col_w, 5, body)
            pdf.ln(2)
        return pdf.get_y()

    left_blocks = [
        (headings.skills, parsed.skills),
        (headings.education, parsed.education),
    ]
    right_blocks = [
        (headings.target, parsed.target),
        (headings.experience, parsed.experience),
    ]
    y_left = write_col(left_x, start_y, left_blocks)
    y_right = write_col(right_x, start_y, right_blocks)
    pdf.set_y(max(y_left, y_right) + 2)
    return pdf


def _pdf_minimal(parsed: ParsedResume, language: str) -> FPDF:
    pdf = _make_pdf(margin=24)
    headings = headings_for(language)
    name = parsed.name or (parsed.raw.splitlines()[0] if parsed.raw else "")
    _pdf_text(pdf, name, size=18, style="", line_h=9)
    if parsed.headline:
        pdf.ln(1)
        _pdf_text(pdf, parsed.headline, size=9, color=_MUTED, line_h=5)
    if not parsed.has_sections:
        pdf.ln(4)
        _pdf_dump(pdf, parsed.raw)
        return pdf

    def section(title: str, body: str) -> None:
        if not body:
            return
        pdf.ln(6)
        pdf.set_text_color(*_MUTED)
        pdf.set_font("DejaVu", size=9)
        pdf.set_x(pdf.l_margin)
        pdf.multi_cell(pdf.epw, 5, title.upper())
        y = pdf.get_y()
        pdf.set_draw_color(203, 213, 225)
        pdf.line(pdf.l_margin, y, pdf.l_margin + pdf.epw, y)
        pdf.ln(3)
        _pdf_text(pdf, body, size=11, line_h=6.5)

    section(headings.target, parsed.target)
    section(headings.experience, parsed.experience)
    section(headings.skills, parsed.skills)
    section(headings.education, parsed.education)
    return pdf


def _shade_cell(cell: object, hex_color: str) -> None:
    tc = cell._tc  # noqa: SLF001
    tc_pr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tc_pr.append(shd)


def _run(paragraph: object, text: str, *, size: int = 11, bold: bool = False, color: RGBColor | None = None) -> None:
    run = paragraph.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(size)
    run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def _docx_section(doc: Document, title: str, body: str, *, size: int = 11, heading_size: int = 12) -> None:
    if not body:
        return
    heading = doc.add_paragraph()
    _run(heading, title, size=heading_size, bold=True)
    for line in body.splitlines() or [""]:
        para = doc.add_paragraph()
        _run(para, line, size=size)


def resume_to_docx(
    resume_text: str,
    *,
    template: str = "classic",
    language: str = "kk",
) -> bytes:
    parsed = parse_resume_text(resume_text)
    if template == "modern":
        doc = _docx_modern(parsed, language)
    elif template == "minimal":
        doc = _docx_minimal(parsed, language)
    else:
        doc = _docx_classic(parsed, language)
    out = BytesIO()
    doc.save(out)
    return out.getvalue()


def _docx_classic(parsed: ParsedResume, language: str) -> Document:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    name = parsed.name or (parsed.raw.splitlines()[0] if parsed.raw else "")
    title = doc.add_paragraph()
    _run(title, name, size=18, bold=True)
    if parsed.headline:
        line = doc.add_paragraph()
        _run(line, parsed.headline, size=10, color=RGBColor(*_MUTED))
    headings = headings_for(language)
    if parsed.has_sections:
        _docx_section(doc, headings.target, parsed.target)
        _docx_section(doc, headings.experience, parsed.experience)
        _docx_section(doc, headings.skills, parsed.skills)
        _docx_section(doc, headings.education, parsed.education)
    else:
        for line in parsed.raw.splitlines():
            para = doc.add_paragraph()
            _run(para, line)
    return doc


def _docx_modern(parsed: ParsedResume, language: str) -> Document:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    header = doc.add_table(rows=1, cols=1)
    cell = header.rows[0].cells[0]
    _shade_cell(cell, "0EA5E9")
    name = parsed.name or (parsed.raw.splitlines()[0] if parsed.raw else "")
    p = cell.paragraphs[0]
    _run(p, name, size=18, bold=True, color=RGBColor(*_WHITE))
    if parsed.headline:
        sub = cell.add_paragraph()
        _run(sub, parsed.headline, size=10, color=RGBColor(*_WHITE))

    if not parsed.has_sections:
        for line in parsed.raw.splitlines():
            para = doc.add_paragraph()
            _run(para, line)
        return doc

    headings = headings_for(language)
    body = doc.add_table(rows=1, cols=2)
    left, right = body.rows[0].cells
    left.width = Cm(6.5)
    right.width = Cm(10.5)

    def fill(cell: object, blocks: list[tuple[str, str]]) -> None:
        first = True
        for title, text in blocks:
            if not text:
                continue
            para = cell.paragraphs[0] if first else cell.add_paragraph()
            first = False
            _run(para, title, size=10, bold=True, color=RGBColor(*_SKY))
            for line in text.splitlines() or [""]:
                body_p = cell.add_paragraph()
                _run(body_p, line, size=10)

    fill(left, [(headings.skills, parsed.skills), (headings.education, parsed.education)])
    fill(right, [(headings.target, parsed.target), (headings.experience, parsed.experience)])
    return doc


def _docx_minimal(parsed: ParsedResume, language: str) -> Document:
    doc = Document()
    for section in doc.sections:
        section.top_margin = Cm(2.4)
        section.bottom_margin = Cm(2.4)
        section.left_margin = Cm(2.4)
        section.right_margin = Cm(2.4)
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    name = parsed.name or (parsed.raw.splitlines()[0] if parsed.raw else "")
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    _run(title, name, size=22, bold=False)
    if parsed.headline:
        line = doc.add_paragraph()
        _run(line, parsed.headline, size=9, color=RGBColor(*_MUTED))
    headings = headings_for(language)
    if parsed.has_sections:
        _docx_section(doc, headings.target.upper(), parsed.target, heading_size=10)
        _docx_section(doc, headings.experience.upper(), parsed.experience, heading_size=10)
        _docx_section(doc, headings.skills.upper(), parsed.skills, heading_size=10)
        _docx_section(doc, headings.education.upper(), parsed.education, heading_size=10)
    else:
        for line in parsed.raw.splitlines():
            para = doc.add_paragraph()
            _run(para, line)
    return doc


def export_bytes(
    resume_text: str,
    fmt: str,
    *,
    template: str = "classic",
    language: str = "kk",
) -> tuple[bytes, str, str]:
    """Return (body, media_type, filename) for pdf|docx."""
    fmt = fmt.lower().strip()
    tmpl: ResumeTemplate = template if template in ("classic", "modern", "minimal") else "classic"
    if fmt == "pdf":
        return (
            resume_to_pdf(resume_text, template=tmpl, language=language),
            "application/pdf",
            _safe_filename("resume", "pdf"),
        )
    if fmt == "docx":
        return (
            resume_to_docx(resume_text, template=tmpl, language=language),
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            _safe_filename("resume", "docx"),
        )
    raise ValueError(f"Unsupported export format: {fmt}")
