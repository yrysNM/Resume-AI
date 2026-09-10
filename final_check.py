# -*- coding: utf-8 -*-
from pathlib import Path
from docx import Document

doc = Document(r"C:\Users\janap\Downloads\НИРМ Ырысбек_толық.docx")
lines = [doc.paragraphs[0].text]
for i, p in enumerate(doc.paragraphs):
    t = (p.text or "").strip()
    if t in ("ҚОРЫТЫНДЫ",) or "4 ТАРАУ" in t and i < 120:
        lines.append(f"{i}: {t[:120]}")
lines.append(f"DOCX size: {Path(r'C:/Users/janap/Downloads/НИРМ Ырысбек_толық.docx').stat().st_size}")
lines.append(f"PDF size: {Path(r'C:/Users/janap/Downloads/НИРМ Ырысбек_толық.pdf').stat().st_size}")
Path(r"C:\Users\janap\Downloads\resume-ai-demo\merged\final_check.txt").write_text("\n---\n".join(lines), encoding="utf-8")
